"""按学号汇总：列顺序与排序验证。"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.util.excel_utils import ExcelExporter


def test_export_results_id_summary_sheet_order_and_headers(tmp_path: Path) -> None:
    exporter = ExcelExporter(tmp_path)
    exporter.export_results(
        [
            {
                "file_name": "A.docx",
                "student_id": "10",
                "student_name": "甲",
                "score": 80.0,
                "status": "成功",
                "comment": "评语A",
                "score_target_max": 100,
                "grader_results": [
                    {
                        "model_index": 1,
                        "status": "成功",
                        "score": 80,
                        "comment": "c1",
                        "sections": [
                            {"name": "内容", "score": 30, "max_score": 40},
                            {"name": "表达", "score": 20, "max_score": 30},
                        ],
                    }
                ],
            },
            {
                "file_name": "B.docx",
                "student_id": "2",
                "student_name": "乙",
                "score": 90.0,
                "status": "成功",
                "comment": "评语B",
                "score_target_max": 100,
                "grader_results": [
                    {
                        "model_index": 1,
                        "status": "成功",
                        "score": 90,
                        "comment": "c2",
                        "sections": [
                            {"name": "内容", "score": 35, "max_score": 40},
                            {"name": "表达", "score": 25, "max_score": 30},
                        ],
                    }
                ],
            },
        ]
    )

    path = tmp_path / "grade_result.xlsx"
    workbook = load_workbook(path)
    ws = workbook["按学号汇总"]

    headers = [cell.value for cell in ws[1]]
    assert headers == ["学号", "姓名", "当前分数", "内容", "表达", "总体评语", "目标总分"]

    # 应按学号升序排序（2 在 10 前）
    assert ws["A2"].value == "2"
    assert ws["A3"].value == "10"

    # 目标总分列在总体评语后
    assert ws["F2"].value == "评语B"
    assert ws["G2"].value == 100
